# -*- coding: utf-8 -*-
"""Создаёт Excel-файл практикума: 4 листа с формулами и диаграммами."""
import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import RadarChart, ScatterChart, BubbleChart, Reference, Series
from openpyxl.chart.series import SeriesLabel
from openpyxl.chart.data_source import StrRef
from openpyxl.formatting.rule import DataBarRule, ColorScaleRule, IconSetRule
from openpyxl.utils import get_column_letter

D = json.load(open('/home/user/kirill-univer/work/data.json', encoding='utf-8'))
poly, ge, bcg = D['poly'], D['ge'], D['bcg']

THIN = Side(style='thin', color='9AA5B1')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD = PatternFill('solid', fgColor='1F4E78')
HEAD2 = PatternFill('solid', fgColor='2E75B6')
ACC = PatternFill('solid', fgColor='DDEBF7')
OURFILL = PatternFill('solid', fgColor='FFF2CC')
GREEN = PatternFill('solid', fgColor='C6EFCE')
YELL = PatternFill('solid', fgColor='FFEB9C')
RED = PatternFill('solid', fgColor='FFC7CE')
WHITEF = Font(color='FFFFFF', bold=True, size=10)
BOLD = Font(bold=True, size=10)
REG = Font(size=10)
WRAP = Alignment(wrap_text=True, vertical='center', horizontal='center')
WRAPL = Alignment(wrap_text=True, vertical='center', horizontal='left')


def style(ws, rng, font=REG, fill=None, align=WRAP, border=True):
    for row in ws[rng]:
        for c in row:
            c.font = font
            if fill:
                c.fill = fill
            c.alignment = align
            if border:
                c.border = BORDER


wb = Workbook()

# ==================== ЛИСТ 1: МНОГОУГОЛЬНИК ====================
ws = wb.active
ws.title = 'Многоугольник'
ws.sheet_view.showGridLines = False
crit = poly['crit']
comp = poly['companies']
ws['A1'] = 'Таблица 27 – Таблица для построения Многоугольника конкурентоспособности ([ФИО студентов], Вариант 7)'
ws['A1'].font = Font(bold=True, size=12, color='1F4E78')
ws.merge_cells('A1:M1')

# header
ws['A3'] = 'Критерии конкурентоспособности'
for i, cn in enumerate(crit):
    ws.cell(row=3, column=2 + i, value=cn)
ws['L3'] = 'U (без учёта дифференциации атрибутов)'
ws['M3'] = 'U детерм (с учётом дифференциации атрибутов)'
style(ws, 'A3:M3', WHITEF, HEAD)

# company score rows 4..8
for r, cn in enumerate(comp):
    row = 4 + r
    ws.cell(row=row, column=1, value=cn)
    for i in range(10):
        ws.cell(row=row, column=2 + i, value=poly['scores'][cn][i])
    ws.cell(row=row, column=12,
            value=f'=SUMPRODUCT(B{row}:K{row},$B$9:$K$9)')
    ws.cell(row=row, column=13,
            value=f'=SUMPRODUCT(B{row}:K{row},$B$12:$K$12)')
    style(ws, f'A{row}:M{row}')
    ws.cell(row=row, column=1).alignment = WRAPL
    if r == 0:
        style(ws, f'A{row}:M{row}', BOLD, OURFILL)
        ws.cell(row=row, column=1).alignment = WRAPL

# row 9 Важность
ws['A9'] = 'Важность'
for i in range(10):
    ws.cell(row=9, column=2 + i, value=poly['importance'][i])
ws['L9'] = '=SUM(B9:K9)'
style(ws, 'A9:M9', BOLD, ACC)
# row 10 Дифференциация
ws['A10'] = 'Дифференциация (СТАНДОТКЛОН.В)'
for i in range(10):
    col = get_column_letter(2 + i)
    ws.cell(row=10, column=2 + i, value=f'=STDEV.S({col}4:{col}8)')
style(ws, 'A10:M10', REG, ACC)
# row 11 C
ws['A11'] = 'C = Важность × Дифференциация'
for i in range(10):
    col = get_column_letter(2 + i)
    ws.cell(row=11, column=2 + i, value=f'={col}9*{col}10')
ws['L11'] = '=SUM(B11:K11)'
style(ws, 'A11:M11', REG, ACC)
# row 12 Детерминация
ws['A12'] = 'Детерминация'
for i in range(10):
    col = get_column_letter(2 + i)
    ws.cell(row=12, column=2 + i, value=f'={col}11/$L$11')
style(ws, 'A12:M12', REG, ACC)
# row 13 Относительная эффективность (к лучшему конкуренту - 1С:ERP, строка 8)
ws['A13'] = 'Относительная эффективность (к 1С:ERP)'
for i in range(10):
    col = get_column_letter(2 + i)
    ws.cell(row=13, column=2 + i, value=f'={col}4/{col}8')
style(ws, 'A13:M13', REG, ACC)

for r in (9, 10, 11, 12, 13):
    ws.cell(row=r, column=1).alignment = WRAPL

# условное форматирование
ws.conditional_formatting.add('B4:K4', DataBarRule(
    start_type='num', start_value=0, end_type='num', end_value=10,
    color='5B9BD5'))
ws.conditional_formatting.add('B4:K4', IconSetRule('3TrafficLights1', 'percent',
                                                   [0, 33, 67]))
ws.conditional_formatting.add('L4:M8', ColorScaleRule(
    start_type='min', start_color='FFC7CE',
    mid_type='percentile', mid_value=50, mid_color='FFEB9C',
    end_type='max', end_color='C6EFCE'))

ws.column_dimensions['A'].width = 30
for i in range(10):
    ws.column_dimensions[get_column_letter(2 + i)].width = 12
ws.column_dimensions['L'].width = 14
ws.column_dimensions['M'].width = 14
for r in range(3, 14):
    ws.row_dimensions[r].height = 46

# Лепестковая диаграмма (многоугольник)
radar = RadarChart()
radar.type = 'marker'
radar.style = 26
radar.title = 'Многоугольник конкурентоспособности ([ФИО студентов], Вариант 7)'
cats = Reference(ws, min_col=2, max_col=11, min_row=3, max_row=3)
for r in range(4, 9):
    s = Series(Reference(ws, min_col=2, max_col=11, min_row=r, max_row=r),
               title_from_data=False)
    s.tx = SeriesLabel(strRef=StrRef(f"'Многоугольник'!A{r}"))
    radar.series.append(s)
radar.set_categories(cats)
radar.width = 19
radar.height = 13
ws.add_chart(radar, 'A16')

# Точечная диаграмма Важность/Эффективность
sc = ScatterChart()
sc.title = 'Матрица «Важность / Эффективность» ([ФИО студентов], Вариант 7)'
sc.x_axis.title = 'Важность критерия'
sc.y_axis.title = 'Полезность атрибута (балл нашей ИС)'
sc.style = 18
xref = Reference(ws, min_col=2, max_col=11, min_row=9, max_row=9)
yref = Reference(ws, min_col=2, max_col=11, min_row=4, max_row=4)
ser = Series(yref, xref, title='Атрибуты ИС НОРБИТ')
ser.marker.symbol = 'circle'
ser.marker.size = 9
ser.graphicalProperties.line.noFill = True
sc.series.append(ser)
sc.width = 19
sc.height = 13
ws.add_chart(sc, 'A44')

# Таблица 29 - План действий
ws['A72'] = 'Таблица 29 – План действий по улучшению конкурентного преимущества'
ws['A72'].font = Font(bold=True, size=11, color='1F4E78')
ws.merge_cells('A72:C72')
ws['A73'] = 'Атрибуты'
ws['B73'] = 'Полезность атрибута (отн. эффективность)'
ws['C73'] = 'План действий по улучшению конкурентного преимущества'
style(ws, 'A73:C73', WHITEF, HEAD)
plan29 = {
    'Функционал': 'Внедрить модуль «SRM.Аттестация»: автосоставление графиков аттестации, оценка надёжности, рейтинг поставщиков — закрывает ключевое отставание по функционалу.',
    'Гибкость и масштабируемость': 'Перевести продукт на микросервисную архитектуру и облачную модель (SaaS), обеспечить горизонтальное масштабирование.',
    'Кастомизация': 'Добавить low-code конструктор регламентов аттестации, настраиваемых критериев и весов оценки надёжности без программирования.',
    'Разнообразие ценообразования': 'Ввести 3 тарифа (базовый, профессиональный, корпоративный) и подписочную SaaS-модель оплаты.',
    'Рекламная активность': 'Запустить контент-маркетинг, SEO, отраслевые вебинары и кейсы; усилить присутствие на ИТ-площадках.',
    'Известность бренда': 'Продвигать суббренд «НОРБИТ SRM», участвовать в профильных выставках и рейтингах SRM-систем.',
    'Дистрибуция': 'Развить партнёрскую сеть внедренцев и реселлеров, разместить решение на маркетплейсе российского ПО.',
}
r = 74
for i, cn in enumerate(crit):
    ws.cell(row=r, column=1, value=cn)
    ws.cell(row=r, column=2, value=f'=ROUND(N4*0,1)' if False else round(poly['rel_eff'][i], 2))
    txt = plan29.get(cn, 'Атрибут выше порогового значения основного конкурента — улучшение не требуется, удерживать достигнутый уровень.')
    ws.cell(row=r, column=3, value=txt)
    style(ws, f'A{r}:C{r}')
    ws.cell(row=r, column=1).alignment = WRAPL
    ws.cell(row=r, column=3).alignment = WRAPL
    if poly['rel_eff'][i] < 1:
        ws.cell(row=r, column=2).fill = RED
    else:
        ws.cell(row=r, column=2).fill = GREEN
    ws.row_dimensions[r].height = 60
    r += 1
ws.column_dimensions['B'].width = 18
ws.column_dimensions['C'].width = 70

# ==================== ЛИСТ 2: GE / McKinsey ====================
ws2 = wb.create_sheet('GE-McKinsey')
ws2.sheet_view.showGridLines = False
seg = ge['segments']


def ge_block(ws, start, title, crit_list, weights, scores_by_seg, label):
    ws.cell(row=start, column=1, value=title).font = Font(bold=True, size=11, color='1F4E78')
    ws.merge_cells(start_row=start, start_column=1, end_row=start, end_column=8)
    h = start + 1
    ws.cell(row=h, column=1, value=label)
    ws.cell(row=h, column=2, value='Вес фактора')
    ws.cell(row=h, column=3, value=f'Оценка: Сегмент 1\n«{seg[0]}»')
    ws.cell(row=h, column=4, value=f'Оценка: Сегмент 2\n«{seg[1]}»')
    ws.cell(row=h, column=5, value=f'Оценка: Сегмент 3\n«{seg[2]}»')
    ws.cell(row=h, column=6, value='Итоговая оценка Сегмент 1')
    ws.cell(row=h, column=7, value='Итоговая оценка Сегмент 2')
    ws.cell(row=h, column=8, value='Итоговая оценка Сегмент 3')
    style(ws, f'A{h}:H{h}', WHITEF, HEAD)
    # сумма весов сверху
    ws.cell(row=h, column=2, value='Вес фактора')
    for i, cn in enumerate(crit_list):
        rr = h + 1 + i
        ws.cell(row=rr, column=1, value=cn)
        ws.cell(row=rr, column=2, value=weights[i])
        for sidx in range(3):
            ws.cell(row=rr, column=3 + sidx, value=scores_by_seg[seg[sidx]][i])
            col_s = get_column_letter(3 + sidx)
            ws.cell(row=rr, column=6 + sidx, value=f'=$B{rr}*{col_s}{rr}')
        style(ws, f'A{rr}:H{rr}')
        ws.cell(row=rr, column=1).alignment = WRAPL
    tot = h + 1 + len(crit_list)
    ws.cell(row=tot, column=1, value='ИТОГО (сумма)')
    ws.cell(row=tot, column=2, value=f'=SUM(B{h+1}:B{tot-1})')
    for sidx in range(3):
        col = get_column_letter(6 + sidx)
        ws.cell(row=tot, column=6 + sidx, value=f'=SUM({col}{h+1}:{col}{tot-1})')
    style(ws, f'A{tot}:H{tot}', BOLD, ACC)
    return tot


ws2['A1'] = 'Матрица GE / McKinsey ([ФИО студентов], Вариант 7)'
ws2['A1'].font = Font(bold=True, size=12, color='1F4E78')
ws2.merge_cells('A1:H1')
t31_end = ge_block(ws2, 3, 'Таблица 31 – Оценка КОНКУРЕНТОСПОСОБНОСТИ',
                   ge['comp_crit'], ge['wc'], ge['comp_sc'],
                   'Критерии конкурентоспособности товара/услуги')
t33_start = t31_end + 3
t33_end = ge_block(ws2, t33_start, 'Таблица 33 – Оценка ПРИВЛЕКАТЕЛЬНОСТИ рынка',
                   ge['attr_crit'], ge['wa'], ge['attr_sc'],
                   'Критерии привлекательности рынка (сегмента)')

# Таблица 34 - матрица GE
m = t33_end + 3
ws2.cell(row=m, column=1, value='Таблица 34 – Матрица General Electric (GE) / McKinsey').font = Font(bold=True, size=11, color='1F4E78')
ws2.merge_cells(start_row=m, start_column=1, end_row=m, end_column=5)
comp_row = t31_end       # строка ИТОГО конкурентоспособности
attr_row = t33_end       # строка ИТОГО привлекательности
# подписи осей
levels = ['Высокая (6,68–10)', 'Средняя (3,34–6,67)', 'Низкая (0–3,33)']
ws2.cell(row=m + 1, column=1, value='Привлекательность рынка ↓ / Конкурентоспособность →')
ws2.cell(row=m + 1, column=3, value='Низкая (0–3,33)')
ws2.cell(row=m + 1, column=4, value='Средняя (3,34–6,67)')
ws2.cell(row=m + 1, column=5, value='Высокая (6,68–10)')
style(ws2, f'A{m+1}:E{m+1}', WHITEF, HEAD2)
cellnum = [['№1', '№2', '№3'], ['№4', '№5', '№6'], ['№7', '№8', '№9']]
for ri in range(3):
    rr = m + 2 + ri
    ws2.cell(row=rr, column=1, value=levels[ri])
    ws2.cell(row=rr, column=1).fill = HEAD2
    ws2.cell(row=rr, column=1).font = WHITEF
    for ci in range(3):
        ws2.cell(row=rr, column=3 + ci, value=cellnum[ri][ci])
    style(ws2, f'A{rr}:E{rr}')
# разместим сегменты по вычисленным уровням
lvl_idx = {'высокая': 0, 'средняя': 1, 'низкая': 2}
for sidx, s in enumerate(seg):
    ri = lvl_idx[ge['attr_level'][s]]
    ci = lvl_idx[ge['comp_level'][s]]
    rr = m + 2 + ri
    cc = 5 - ci
    cell = ws2.cell(row=rr, column=cc)
    cell.value = (cell.value or '') + f'\nСегмент {sidx+1}: «{s}»\n(привл.={ge["attr_total"][s]:.2f}; конкур.={ge["comp_total"][s]:.2f})'
    cell.fill = GREEN
    cell.alignment = WRAP
for rr in range(m + 2, m + 5):
    ws2.row_dimensions[rr].height = 90
ws2.column_dimensions['A'].width = 34
for col in 'BCDEFGH':
    ws2.column_dimensions[col].width = 19

# ==================== ЛИСТ 3: BCG ====================
ws3 = wb.create_sheet('BCG')
ws3.sheet_view.showGridLines = False
ws3['A1'] = 'Матрица BCG (БКГ) ([ФИО студентов], Вариант 7)'
ws3['A1'].font = Font(bold=True, size=12, color='1F4E78')
ws3.merge_cells('A1:L1')
ws3['A3'] = 'Таблица 36 – Расчёт темпа роста рынка и доли рынка'
ws3['A3'].font = Font(bold=True, size=11, color='1F4E78')
ws3.merge_cells('A3:L3')
hdr = ['Название группы', 'Объём продаж 2024, руб.', 'Объём продаж 2025, руб.',
       'Темп роста рынка, %', 'Ёмкость рынка', 'Взвешенный темп роста, %',
       'Рост для матрицы BCG', 'Доля рынка бренда', 'Доля рынка ключ. конкурента',
       'Относительная доля рынка', 'Доля для матрицы BCG', 'Название продукта']
for i, h in enumerate(hdr):
    ws3.cell(row=4, column=1 + i, value=h)
style(ws3, 'A4:L4', WHITEF, HEAD)
rows = bcg['rows']
for r, row in enumerate(rows):
    rr = 5 + r
    ws3.cell(row=rr, column=1, value=f'Товар {r+1}')
    ws3.cell(row=rr, column=2, value=row['s24'])
    ws3.cell(row=rr, column=3, value=row['s25'])
    ws3.cell(row=rr, column=4, value=f'=C{rr}/B{rr}*100')
    ws3.cell(row=rr, column=5, value=row['cap'])
    ws3.cell(row=rr, column=6, value=f'=E{rr}*D{rr}/$E$11')
    ws3.cell(row=rr, column=7, value=f'=IF(F{rr}>10,"высокий","низкий")')
    ws3.cell(row=rr, column=8, value=row['share_brand'] / 100)
    ws3.cell(row=rr, column=9, value=row['share_comp'] / 100)
    ws3.cell(row=rr, column=10, value=f'=H{rr}/I{rr}')
    ws3.cell(row=rr, column=11, value=f'=IF(J{rr}>1,"высокая","низкая")')
    ws3.cell(row=rr, column=12, value=row['name'])
    ws3.cell(row=rr, column=8).number_format = '0%'
    ws3.cell(row=rr, column=9).number_format = '0%'
    ws3.cell(row=rr, column=4).number_format = '0.0'
    ws3.cell(row=rr, column=6).number_format = '0.00'
    ws3.cell(row=rr, column=10).number_format = '0.00'
    style(ws3, f'A{rr}:L{rr}')
    ws3.cell(row=rr, column=12).alignment = WRAPL
ws3['A11'] = 'ИТОГО'
ws3['B11'] = '=SUM(B5:B10)'
ws3['C11'] = '=SUM(C5:C10)'
ws3['E11'] = '=SUM(E5:E10)'
style(ws3, 'A11:L11', BOLD, ACC)
ws3.conditional_formatting.add('F5:F10', DataBarRule(
    start_type='min', end_type='max', color='70AD47'))
ws3.conditional_formatting.add('J5:J10', DataBarRule(
    start_type='min', end_type='max', color='5B9BD5'))
ws3.column_dimensions['A'].width = 14
ws3.column_dimensions['L'].width = 34
for col in 'BCDEFGHIJK':
    ws3.column_dimensions[col].width = 14
for rr in range(4, 12):
    ws3.row_dimensions[rr].height = 42

# Пузырьковая диаграмма
bub = BubbleChart()
bub.title = 'Матрица BCG «Темп роста – Доля рынка» ([ФИО студентов], Вариант 7)'
bub.x_axis.title = 'Относительная доля рынка'
bub.y_axis.title = 'Взвешенный темп роста рынка, %'
bub.style = 18
xv = Reference(ws3, min_col=10, min_row=5, max_row=10)
yv = Reference(ws3, min_col=6, min_row=5, max_row=10)
zv = Reference(ws3, min_col=3, min_row=5, max_row=10)
bub.series.append(Series(yv, xv, zv, title='Товары компании'))
bub.width = 20
bub.height = 13
ws3.add_chart(bub, 'A14')

# Таблица 37 - матрица BCG
m = 42
ws3.cell(row=m, column=1, value='Таблица 37 – Матрица БКГ').font = Font(bold=True, size=11, color='1F4E78')
ws3.merge_cells(start_row=m, start_column=1, end_row=m, end_column=5)
quads = {'ТРУДНЫЕ ДЕТИ': [], 'ЗВЕЗДЫ': [], 'СОБАКИ': [], 'ДОЙНЫЕ КОРОВЫ': []}
for r, row in enumerate(rows):
    quads[row['quad']].append((f'Товар {r+1}', row['name'], row['s25']))


def quad_cell(ws, r, c, title, items):
    ws.cell(row=r, column=c, value=title).font = WHITEF
    ws.cell(row=r, column=c).fill = HEAD2
    txt = title + ':\n' + ('\n'.join(f'{a} {b} — {z} руб.' for a, b, z in items)
                           if items else '—')
    cc = ws.cell(row=r + 1, column=c, value=txt)
    cc.alignment = WRAP
    ws.merge_cells(start_row=r + 1, start_column=c, end_row=r + 1, end_column=c + 1)
    ws.merge_cells(start_row=r, start_column=c, end_row=r, end_column=c + 1)


ws3.cell(row=m + 1, column=1, value='Темп роста ВЫСОКИЙ (>10%)').font = BOLD
quad_cell(ws3, m + 2, 2, 'ТРУДНЫЕ ДЕТИ', quads['ТРУДНЫЕ ДЕТИ'])
quad_cell(ws3, m + 2, 4, 'ЗВЁЗДЫ', quads['ЗВЕЗДЫ'])
ws3.cell(row=m + 4, column=1, value='Темп роста НИЗКИЙ (<10%)').font = BOLD
quad_cell(ws3, m + 5, 2, 'СОБАКИ', quads['СОБАКИ'])
quad_cell(ws3, m + 5, 4, 'ДОЙНЫЕ КОРОВЫ', quads['ДОЙНЫЕ КОРОВЫ'])
ws3.cell(row=m + 7, column=2, value='Относительная доля рынка НИЗКАЯ (<1)').font = BOLD
ws3.cell(row=m + 7, column=4, value='Относительная доля рынка ВЫСОКАЯ (>1)').font = BOLD
for rr in (m + 3, m + 6):
    ws3.row_dimensions[rr].height = 70
for rng in (f'B{m+2}:E{m+3}', f'B{m+5}:E{m+6}'):
    style(ws3, rng)

# ==================== ЛИСТ 4: МАТРИЦА АНСОФФА ====================
ws4 = wb.create_sheet('Матрица Ансоффа')
ws4.sheet_view.showGridLines = False
ws4['A1'] = 'Матрица Ансоффа ([ФИО студентов], Вариант 7)'
ws4['A1'].font = Font(bold=True, size=12, color='1F4E78')
ws4.merge_cells('A1:E1')

ansoff = json.load(open('/home/user/kirill-univer/work/ansoff.json', encoding='utf-8'))


def ansoff_table(ws, start, blk):
    ws.cell(row=start, column=1, value=blk['title']).font = Font(bold=True, size=11, color='1F4E78')
    ws.merge_cells(start_row=start, start_column=1, end_row=start, end_column=5)
    q = start + 1
    ws.cell(row=q, column=1, value=blk['question'])
    ws.merge_cells(start_row=q, start_column=1, end_row=q, end_column=5)
    style(ws, f'A{q}:E{q}', WHITEF, HEAD)
    ws.cell(row=q, column=1).alignment = WRAPL
    h = q + 1
    ws.cell(row=h, column=1, value='Параметр оценки')
    ws.cell(row=h, column=2, value='Возможна')
    ws.cell(row=h, column=3, value='Вероятна')
    ws.cell(row=h, column=4, value='Не возможна')
    ws.cell(row=h, column=5, value='Варианты ответа')
    style(ws, f'A{h}:E{h}', WHITEF, HEAD2)
    d = h + 1
    ws.cell(row=d, column=1, value='Описание рынка и товара')
    ws.cell(row=d, column=2, value=blk['describe'])
    ws.merge_cells(start_row=d, start_column=2, end_row=d, end_column=5)
    style(ws, f'A{d}:E{d}', BOLD, ACC)
    ws.cell(row=d, column=2).alignment = WRAPL
    ws.row_dimensions[d].height = 54
    rr = d + 1
    for param, mark, chosen, variants in blk['rows']:
        ws.cell(row=rr, column=1, value=param)
        col = {'green': 2, 'yellow': 3, 'red': 4}[mark]
        fill = {'green': GREEN, 'yellow': YELL, 'red': RED}[mark]
        cell = ws.cell(row=rr, column=col, value=chosen)
        ws.cell(row=rr, column=5, value=variants)
        style(ws, f'A{rr}:E{rr}')
        cell.fill = fill
        ws.cell(row=rr, column=1).alignment = WRAPL
        ws.cell(row=rr, column=5).alignment = WRAPL
        ws.cell(row=rr, column=col).alignment = WRAPL
        ws.row_dimensions[rr].height = 48
        rr += 1
    ws.cell(row=rr, column=1, value='ВЫВОД ПО СТРАТЕГИИ:')
    ws.cell(row=rr, column=2, value=blk['verdict'])
    ws.merge_cells(start_row=rr, start_column=2, end_row=rr, end_column=5)
    style(ws, f'A{rr}:E{rr}', BOLD, ACC)
    ws.cell(row=rr, column=2).alignment = WRAPL
    ws.row_dimensions[rr].height = 60
    return rr


pos = 3
for blk in ansoff['tables']:
    pos = ansoff_table(ws4, pos, blk) + 3

# Таблица 43 - матрица Ансоффа
ws4.cell(row=pos, column=1, value='Таблица 43 – Матрица Ансоффа').font = Font(bold=True, size=11, color='1F4E78')
ws4.merge_cells(start_row=pos, start_column=1, end_row=pos, end_column=4)
ws4.cell(row=pos + 1, column=3, value='Существующий продукт')
ws4.cell(row=pos + 1, column=4, value='Новый продукт')
ws4.cell(row=pos + 2, column=1, value='Существующий рынок')
ws4.cell(row=pos + 3, column=1, value='Новый рынок')
style(ws4, f'C{pos+1}:D{pos+1}', WHITEF, HEAD2)
style(ws4, f'A{pos+2}:A{pos+3}', WHITEF, HEAD2)
m43 = ansoff['matrix43']
ws4.cell(row=pos + 2, column=3, value=m43['penetration'])
ws4.cell(row=pos + 2, column=4, value=m43['product'])
ws4.cell(row=pos + 3, column=3, value=m43['market'])
ws4.cell(row=pos + 3, column=4, value=m43['diversification'])
style(ws4, f'C{pos+2}:D{pos+3}')
for rr in (pos + 2, pos + 3):
    ws4.row_dimensions[rr].height = 110
    for cc in (3, 4):
        ws4.cell(row=rr, column=cc).alignment = WRAP

# Таблица 44 - возможности
p44 = pos + 6
ws4.cell(row=p44, column=1, value='Таблица 44 – Возможности реализации каждой стратегии').font = Font(bold=True, size=11, color='1F4E78')
ws4.merge_cells(start_row=p44, start_column=1, end_row=p44, end_column=4)
for i, h in enumerate(['Вариант стратегии', 'Возможность', 'Описание', 'Ключевые источники роста']):
    ws4.cell(row=p44 + 1, column=1 + i, value=h)
style(ws4, f'A{p44+1}:D{p44+1}', WHITEF, HEAD)
rr = p44 + 2
for row in ansoff['table44']:
    for i, v in enumerate(row):
        ws4.cell(row=rr, column=1 + i, value=v)
    style(ws4, f'A{rr}:D{rr}')
    for cc in range(1, 5):
        ws4.cell(row=rr, column=cc).alignment = WRAPL
    ws4.row_dimensions[rr].height = 120
    rr += 1
ws4.column_dimensions['A'].width = 30
for col in 'BCDE':
    ws4.column_dimensions[col].width = 32

wb.save('/home/user/kirill-univer/Практикум МИТиУ_Вариант 7.xlsx')
print('Excel сохранён OK')
