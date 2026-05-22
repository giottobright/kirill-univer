# -*- coding: utf-8 -*-
"""Заполнение практикума МИТиУ (Вариант 7): таблицы, абзацы, диаграммы."""
import copy, json
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.table import Table
from docx.text.paragraph import Paragraph
import content as C

DATA = json.load(open('/home/user/kirill-univer/work/data.json', encoding='utf-8'))
ANSOFF = json.load(open('/home/user/kirill-univer/work/ansoff.json', encoding='utf-8'))
poly, ge, bcg = DATA['poly'], DATA['ge'], DATA['bcg']
A = '/home/user/kirill-univer/work/assets/'

SRC = '/home/user/kirill-univer/2026_Практикум_МИТиУ.docx'
OUT = '/home/user/kirill-univer/Практикум МИТиУ_Вариант 7.docx'
doc = docx.Document(SRC)
P = doc.paragraphs
T = doc.tables

NAVY = RGBColor(0x1F, 0x4E, 0x78)
GREEN = RGBColor(0x2E, 0x6B, 0x16)
BLUEC = RGBColor(0x1F, 0x4E, 0x78)
ORANGEC = RGBColor(0xB0, 0x4A, 0x00)
PURPLE = RGBColor(0x5B, 0x2C, 0x83)
SH_GREEN, SH_YELL, SH_RED = 'C6EFCE', 'FFEB9C', 'FFC7CE'


def shade(cell, hex6):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:fill'), hex6)
    tcPr.append(shd)


def clear_para(p):
    for r in list(p.runs):
        r._element.getparent().remove(r._element)


def fill_cell(cell, text, size=9, bold=False, color=None, center=False):
    cell.text = ''
    lines = str(text).split('\n')
    for i, ln in enumerate(lines):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(ln)
        run.font.size = Pt(size)
        run.font.name = 'Times New Roman'
        run.bold = bold
        if color:
            run.font.color.rgb = color


def fill_cell_colored(cell, parts, size=9):
    """parts: список (text, RGBColor|None)."""
    cell.text = ''
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for text, col in parts:
        run = p.add_run(text)
        run.font.size = Pt(size)
        run.font.name = 'Times New Roman'
        if col:
            run.font.color.rgb = col
            run.bold = True


def set_answer(para, prefix, text, prefix_bold=True):
    clear_para(para)
    if prefix:
        r = para.add_run(prefix + ' ')
        r.bold = prefix_bold
        r.font.size = Pt(11)
        r.font.name = 'Times New Roman'
    r2 = para.add_run(text)
    r2.font.size = Pt(11)
    r2.font.name = 'Times New Roman'


def insert_paragraph_after(para, text='', bold=False, italic=False, size=11,
                           align='left', color=None):
    new_p = OxmlElement('w:p')
    para._p.addnext(new_p)
    np = Paragraph(new_p, para._parent)
    np.alignment = {'left': WD_ALIGN_PARAGRAPH.LEFT, 'center': WD_ALIGN_PARAGRAPH.CENTER,
                    'justify': WD_ALIGN_PARAGRAPH.JUSTIFY}[align]
    if text:
        run = np.add_run(text)
        run.bold = bold
        run.italic = italic
        run.font.size = Pt(size)
        run.font.name = 'Times New Roman'
        if color:
            run.font.color.rgb = color
    return np


def insert_image_after(para, img, width=6.3, caption=None):
    np = insert_paragraph_after(para, '', align='center')
    np.add_run().add_picture(img, width=Inches(width))
    last = np
    if caption:
        last = insert_paragraph_after(np, caption, italic=True, size=9,
                                      align='center', color=NAVY)
    return last


def fill_table(idx, layout):
    """layout: dict (row,col)->text для простого заполнения."""
    tbl = T[idx]
    for (r, c), text in layout.items():
        fill_cell(tbl.rows[r].cells[c], text)


# ================= ТИТУЛ =================
fill_cell(T[0].rows[0].cells[0],
          'ФИО студентов, выполняющих работу: [ФИО студента], '
          '[Член команды 2], [Член команды 3]          Вариант № 7'
          '          Группа: [Группа]', size=11, bold=True, center=True)

# ================= ТАБЛИЦА 2 =================
for c in range(6):
    fill_cell(T[2].rows[1].cells[c], C.T2[c], size=9,
              center=(c == 0))

# ================= ТАБЛИЦА 3 =================
for i in range(15):
    cell = T[3].rows[i + 1].cells[1]
    if C.T3[i] == 'IDEA4':
        cols = [BLUEC, GREEN, ORANGEC, PURPLE]
        fill_cell_colored(cell, [(C.IDEA4[k], cols[k]) for k in range(4)])
    else:
        fill_cell(cell, C.T3[i])

# ================= ТАБЛИЦА 4: SWOT =================
fill_cell(T[4].rows[1].cells[0], C.SWOT['S'])
fill_cell(T[4].rows[1].cells[1], C.SWOT['W'])
fill_cell(T[4].rows[3].cells[0], C.SWOT['O'])
fill_cell(T[4].rows[3].cells[1], C.SWOT['T'])

# ================= ТАБЛИЦА 5 =================
for i in range(3):
    fill_cell(T[5].rows[i + 1].cells[1], C.T5[i][0])
    fill_cell(T[5].rows[i + 1].cells[2], C.T5[i][1], center=True)

# ================= ТАБЛИЦА 6 =================
for i in range(3):
    fill_cell(T[6].rows[i].cells[1], C.T6[i])

# ================= ТАБЛИЦА 7 =================
for i in range(9):
    fill_cell(T[7].rows[i].cells[1], C.T7[i])

# ================= ТАБЛИЦА 8 =================
for i in range(6):
    fill_cell(T[8].rows[i].cells[1], C.T8[i])

# ================= ТАБЛИЦА 9 (CJM) =================
t9 = T[9]
for ci in range(6):
    fill_cell(t9.rows[0].cells[ci + 1], 'Ситуация ' + str(ci + 1) + '. ' +
              C.CJM_STAGES[ci], size=8, bold=True, center=True)
rows9 = [C.CJM_ACTIONS, C.CJM_MOODS, C.CJM_TOUCH, C.CJM_PAINS, C.CJM_OPP]
for ri, vals in enumerate(rows9):
    for ci in range(6):
        fill_cell(t9.rows[ri + 1].cells[ci + 1], vals[ci], size=8,
                  center=(ri == 1))

# ================= ТАБЛИЦА 10 =================
fill_cell(T[10].rows[1].cells[0], C.T10[0])
fill_cell(T[10].rows[1].cells[1], C.T10[1])

# ================= ТАБЛИЦА 11 =================
for i in range(3):
    fill_cell(T[11].rows[i].cells[1], C.T11[i])

# ================= ТАБЛИЦА 14 =================
for c in range(3):
    fill_cell(T[14].rows[1].cells[c], C.T14[c])

# ================= ТАБЛИЦА 16 (выбор комм. стратегии) =================
for c in range(3):
    fill_cell(T[17].rows[1].cells[c], C.T16[c])

# ================= ТАБЛИЦА 17 (инструменты продвижения) =================
for i in range(9):
    fill_cell(T[18].rows[i + 1].cells[1], C.T17[i][0])
    fill_cell(T[18].rows[i + 1].cells[2], C.T17[i][1], center=True)

# ================= ТАБЛИЦА 18 (4P) =================
for i in range(4):
    fill_cell(T[19].rows[i].cells[2], C.T18[i])

# ================= ТАБЛИЦА 19 (Impact mapping) =================
for i in range(5):
    fill_cell(T[20].rows[i].cells[1], C.T19[i])

# ================= ТАБЛИЦА 20 (User Story) =================
t21 = T[21]
cols3 = [BLUEC, GREEN, ORANGEC]
while len(t21.rows) < len(C.US) + 1:
    t21.add_row()
for i, us in enumerate(C.US):
    row = t21.rows[i + 1]
    fill_cell(row.cells[1], us[0], size=9)
    fill_cell(row.cells[2], us[1], size=9)
    fill_cell_colored(row.cells[3],
                      [(us[2] + ' ', cols3[0]), (us[3] + ', ', cols3[1]),
                       (us[4] + '.', cols3[2])], size=9)

# ================= ТАБЛИЦА 21 (Given/When/Then) =================
def fill_gwt(tbl, g):
    fill_cell(tbl.rows[0].cells[0], 'User Story: ' + g['us'], size=9, bold=True)
    fill_cell(tbl.rows[0].cells[1], 'User Story: ' + g['us'], size=9, bold=True)
    fill_cell(tbl.rows[1].cells[1], g['Сценарий'])
    fill_cell(tbl.rows[2].cells[1], 'Дано: ' + g['Дано'])
    fill_cell(tbl.rows[3].cells[1], 'Когда: ' + g['Когда'])
    fill_cell(tbl.rows[4].cells[1], 'И ' + g['И1'])
    fill_cell(tbl.rows[5].cells[1], 'Тогда: ' + g['Тогда'])
    fill_cell(tbl.rows[6].cells[1], 'И ' + g['И2'])


fill_gwt(T[22], C.GWT[0])

# ================= ТАБЛИЦА 21-прецеденты (Use Case) =================
t23 = T[23]
for i, (us, uc) in enumerate(C.T23):
    fill_cell(t23.rows[i + 1].cells[0], str(i + 1), center=True)
    fill_cell(t23.rows[i + 1].cells[1], us, size=9)
    fill_cell(t23.rows[i + 1].cells[2], uc, size=9)

# ================= ТАБЛИЦА 22 (Вариант использования) =================
for i in range(9):
    fill_cell(T[24].rows[i].cells[1], C.T22UC[i])

# ================= ТАБЛИЦА 23 практикума (ключевые фразы) =================
t25 = T[25]
for i, kw in enumerate(C.KEYWORDS):
    row = t25.rows[i + 1]
    fill_cell(row.cells[0], str(i + 1), center=True)
    for c in range(4):
        fill_cell(row.cells[c + 1], kw[c], center=True)

# ================= ТАБЛИЦА 24 (Яндекс.Директ) =================
for i in range(6):
    fill_cell(T[26].rows[i].cells[1], C.DIRECT[i])

# ================= ТАБЛИЦА 27 практикума (текст для Advego) =================
fill_cell(T[27].rows[0].cells[0], C.ANSWERS['advego_text'], size=10)

# ================= ТАБЛИЦА 25 практикума (запуск) =================
for i in range(5):
    fill_cell(T[28].rows[i].cells[1], C.LAUNCH[i])

# ================= ТАБЛИЦА 26 (КП) =================
for i in range(8):
    fill_cell(T[29].rows[i].cells[2], C.KP[i])

print('Текстовые таблицы заполнены')


# ===================== АНАЛИТИЧЕСКИЕ ТАБЛИЦЫ =====================
def n(x, d=2):
    return f'{x:.{d}f}'.replace('.', ',')


def sh_by_val(cell, v):
    shade(cell, SH_GREEN if v >= 8 else SH_YELL if v >= 4 else SH_RED)


# ----- ТАБЛИЦА 27: Многоугольник -----
t30 = T[30]
comp_names = poly['companies']
for ri, cn in enumerate(comp_names):
    row = t30.rows[2 + ri]
    fill_cell(row.cells[0], cn, size=8, bold=(ri == 0))
    for ci in range(10):
        fill_cell(row.cells[1 + ci], str(poly['scores'][cn][ci]), size=9, center=True)
        if ri == 0:
            sh_by_val(row.cells[1 + ci], poly['scores'][cn][ci])
    fill_cell(row.cells[11], n(poly['U'][cn]), size=9, center=True)
    fill_cell(row.cells[12], n(poly['Udet'][cn]), size=9, center=True)
for ci in range(10):
    fill_cell(t30.rows[7].cells[1 + ci], n(poly['importance'][ci]), size=9, center=True)
    fill_cell(t30.rows[8].cells[1 + ci], n(poly['diff'][ci]), size=9, center=True)
    fill_cell(t30.rows[9].cells[1 + ci], n(poly['C'][ci], 3), size=9, center=True)
    fill_cell(t30.rows[10].cells[1 + ci], n(poly['determ'][ci], 3), size=9, center=True)
    fill_cell(t30.rows[11].cells[1 + ci], n(poly['rel_eff'][ci]), size=9, center=True)
fill_cell(t30.rows[7].cells[11], '1,00', size=9, center=True)

# ----- ТАБЛИЦА 29: План действий -----
t34 = T[34]
plan29 = {
    'Размер бизнеса заказчика': 'Атрибут выше уровня основного конкурента — удерживать достигнутый уровень.',
    'Функционал': 'Внедрить модуль «SRM.Аттестация»: автосоставление графиков аттестации, оценка надёжности и рейтинг поставщиков — закрывает ключевое отставание.',
    'Гибкость и масштабируемость': 'Перевести продукт на микросервисную архитектуру и облачную SaaS-модель, обеспечить горизонтальное масштабирование.',
    'Производительность платформы': 'Атрибут выше уровня основного конкурента — удерживать достигнутый уровень.',
    'Кастомизация': 'Добавить low-code конструктор регламентов аттестации и настраиваемых критериев оценки надёжности без программирования.',
    'Универсальность решения': 'Атрибут выше уровня основного конкурента — удерживать достигнутый уровень.',
    'Разнообразие ценообразования': 'Ввести 3 тарифа (базовый, профессиональный, корпоративный) и подписочную SaaS-модель оплаты.',
    'Рекламная активность': 'Запустить контент-маркетинг, SEO, отраслевые вебинары и кейсы, усилить присутствие на ИТ-площадках.',
    'Известность бренда': 'Продвигать суббренд «НОРБИТ SRM», участвовать в профильных выставках и рейтингах SRM-систем.',
    'Дистрибуция': 'Развить партнёрскую сеть внедренцев и реселлеров, разместить решение на маркетплейсе российского ПО.',
}
for i, crit in enumerate(poly['crit']):
    row = t34.rows[i + 1]
    re = poly['rel_eff'][i]
    fill_cell(row.cells[1], n(re) + (' (выше конкурента)' if re >= 1 else ' (ниже конкурента)'),
              size=9, center=True)
    shade(row.cells[1], SH_GREEN if re >= 1 else SH_RED)
    fill_cell(row.cells[2], plan29[crit], size=9)

# ----- ТАБЛИЦЫ 31 и 33: GE/McKinsey -----
def fill_ge(tbl, crit_list, weights, scores, totals):
    fill_cell(tbl.rows[1].cells[1], '100%', center=True, bold=True)
    for s in range(3):
        fill_cell(tbl.rows[1].cells[2 + s],
                  f'Сегмент {s+1}: «{C.SEGMENTS[s]}»', size=8, bold=True, center=True)
        fill_cell(tbl.rows[1].cells[5 + s], n(totals[C.SEGMENTS[s]]), bold=True, center=True)
    for i, cr in enumerate(crit_list):
        row = tbl.rows[2 + i]
        fill_cell(row.cells[0], cr, size=8)
        fill_cell(row.cells[1], n(weights[i]), center=True)
        for s in range(3):
            sc = scores[C.SEGMENTS[s]][i]
            fill_cell(row.cells[2 + s], str(sc), center=True)
            fill_cell(row.cells[5 + s], n(weights[i] * sc), center=True)


fill_ge(T[36], ge['comp_crit'], ge['wc'], ge['comp_sc'], ge['comp_total'])
fill_ge(T[38], ge['attr_crit'], ge['wa'], ge['attr_sc'], ge['attr_total'])

# ----- ТАБЛИЦА 34: матрица GE -----
t39 = T[39]
lvl_row = {'высокая': 0, 'средняя': 1, 'низкая': 2}
lvl_col = {'низкая': 2, 'средняя': 3, 'высокая': 4}
ge_cellnum = {(0, 3): '№2 — стратегия победителя, усиление конкурентных преимуществ',
              (1, 3): '№5 — средние позиции, осторожное продолжение бизнеса',
              (1, 4): '№6 — стратегия победителя, рост и максимизация выгоды'}
for si, seg in enumerate(ge['segments']):
    r = lvl_row[ge['attr_level'][seg]]
    c = lvl_col[ge['comp_level'][seg]]
    note = ge_cellnum.get((r, c), '')
    fill_cell(t39.rows[r].cells[c],
              f'Сегмент {si+1}: «{seg}»\n(привл. {n(ge["attr_total"][seg])}; '
              f'конкур. {n(ge["comp_total"][seg])})\n{note}', size=8, bold=True)
    shade(t39.rows[r].cells[c], SH_GREEN)

# ----- ТАБЛИЦА 36: расчёт BCG -----
t41 = T[41]
rows_b = bcg['rows']
for i, rb in enumerate(rows_b):
    row = t41.rows[3 + i]
    vals = [f'Товар {i+1}', str(rb['s24']), str(rb['s25']), n(rb['tr'], 1) + '%',
            str(rb['cap']), n(rb['trw'], 2) + '%', rb['growth'],
            str(rb['share_brand']) + '%', str(rb['share_comp']) + '%',
            n(rb['odr']), rb['odr_txt'], rb['name']]
    for ci, v in enumerate(vals):
        fill_cell(row.cells[ci], v, size=8, center=(0 < ci < 11))
fill_cell(t41.rows[9].cells[0], 'ИТОГО', bold=True, center=True)
fill_cell(t41.rows[9].cells[1], str(sum(r['s24'] for r in rows_b)), bold=True, center=True)
fill_cell(t41.rows[9].cells[2], str(sum(r['s25'] for r in rows_b)), bold=True, center=True)
fill_cell(t41.rows[9].cells[4], str(bcg['sumE']), bold=True, center=True)

# ----- ТАБЛИЦА 37: матрица BCG -----
t42 = T[42]
quads = {'ТРУДНЫЕ ДЕТИ': [], 'ЗВЕЗДЫ': [], 'СОБАКИ': [], 'ДОЙНЫЕ КОРОВЫ': []}
for i, rb in enumerate(rows_b):
    quads[rb['quad']].append((f'Товар {i+1} {rb["name"]}', rb['s25']))


def put_quad(rows_idx, c_name, c_val, items):
    for k in range(len(rows_idx)):
        if k < len(items):
            chunk = items[k::len(rows_idx)] if False else None
    # распределяем: по одному на строку, излишек — в последнюю
    per = [[] for _ in rows_idx]
    for j, it in enumerate(items):
        per[min(j, len(rows_idx) - 1)].append(it)
    for k, ri in enumerate(rows_idx):
        names = '\n'.join(it[0] for it in per[k])
        vals = '\n'.join(str(it[1]) for it in per[k])
        fill_cell(t42.rows[ri].cells[c_name], names, size=8)
        fill_cell(t42.rows[ri].cells[c_val], str(vals), size=8, center=True)


put_quad([2, 3], 2, 3, quads['ТРУДНЫЕ ДЕТИ'])
put_quad([2, 3], 4, 5, quads['ЗВЕЗДЫ'])
put_quad([7, 8], 2, 3, quads['СОБАКИ'] or [('—', '—')])
put_quad([7, 8], 4, 5, quads['ДОЙНЫЕ КОРОВЫ'])
fill_cell(t42.rows[4].cells[3], str(sum(x[1] for x in quads['ТРУДНЫЕ ДЕТИ'])), bold=True, center=True)
fill_cell(t42.rows[4].cells[5], str(sum(x[1] for x in quads['ЗВЕЗДЫ'])), bold=True, center=True)
fill_cell(t42.rows[9].cells[3], str(sum(x[1] for x in quads['СОБАКИ'])), bold=True, center=True)
fill_cell(t42.rows[9].cells[5], str(sum(x[1] for x in quads['ДОЙНЫЕ КОРОВЫ'])), bold=True, center=True)

# ----- ТАБЛИЦА 38: выводы BCG -----
t43 = T[43]
fill_cell(t43.rows[2].cells[0],
          'Товары 2, 4, 6 («Управление автотранспортом», «Управление складом (WMS)», '
          '«Сервис Деск») — высокий темп роста рынка при низкой относительной доле. '
          'Стратегия: селективное инвестирование в наиболее перспективные товары, '
          'превращение лидеров в «звёзды»; от бесперспективных — отказ.', size=9)
fill_cell(t43.rows[2].cells[1],
          'Товары 3, 5 («НОРБИТ: ТОИР», «Управление маркетингом») — высокий рост и '
          'высокая относительная доля рынка. Стратегия: сохранение лидерства, защита '
          'позиций, реинвестирование прибыли в развитие.', size=9)
fill_cell(t43.rows[4].cells[0],
          'Товаров-«собак» в портфеле нет — благоприятный признак: ресурсы не '
          'расходуются на бесперспективные продукты.', size=9)
fill_cell(t43.rows[4].cells[1],
          'Товар 1 («НОРБИТ: Управление закупками (SRM)») — низкий темп роста при '
          'высокой доле рынка. Стратегия: получение максимальной прибыли при '
          'минимальных вложениях; средства направляются на развитие «звёзд» и '
          '«трудных детей».', size=9)
fill_cell(t43.rows[5].cells[1],
          'Портфель сбалансирован достаточно: «дойная корова» (Т1) генерирует '
          'денежный поток, «звёзды» (Т3, Т5) обеспечивают рост, «трудные дети» '
          '(Т2, Т4, Т6) — будущий потенциал. Риск — единственная «дойная корова»: '
          'денежный поток зависит от одного товара. Новый модуль «SRM.Аттестация» '
          'призван оживить рост SRM-направления и удержать его в роли источника '
          'прибыли.', size=9)

# ----- ТАБЛИЦЫ 39-42: Ансофф -----
ansoff_doc_idx = [44, 45, 46, 47]
mark_col = {'green': 1, 'yellow': 2, 'red': 3}
mark_sh = {'green': SH_GREEN, 'yellow': SH_YELL, 'red': SH_RED}
for bi, blk in enumerate(ANSOFF['tables']):
    tbl = T[ansoff_doc_idx[bi]]
    fill_cell(tbl.rows[2].cells[1], blk['describe'], size=8)
    for pi, (param, mark, chosen, variants) in enumerate(blk['rows']):
        row = tbl.rows[3 + pi]
        cidx = mark_col[mark]
        fill_cell(row.cells[cidx], chosen, size=8, center=True)
        shade(row.cells[cidx], mark_sh[mark])
    vp = OxmlElement('w:p')
    tbl._element.addnext(vp)
    vpar = Paragraph(vp, tbl._parent)
    rr = vpar.add_run('Вывод по стратегии: ' + blk['verdict'])
    rr.bold = True
    rr.font.size = Pt(10)
    rr.font.name = 'Times New Roman'

# ----- ТАБЛИЦА 43: матрица Ансоффа -----
t48 = T[48]
m43 = ANSOFF['matrix43']
fill_cell(t48.rows[2].cells[2], m43['penetration'], size=9)
fill_cell(t48.rows[2].cells[3], m43['product'], size=9)
shade(t48.rows[2].cells[3], SH_GREEN)
fill_cell(t48.rows[3].cells[2], m43['market'], size=9)
fill_cell(t48.rows[3].cells[3], m43['diversification'], size=9)
shade(t48.rows[3].cells[3], SH_RED)

# ----- ТАБЛИЦА 44: возможности реализации -----
t49 = T[49]
for i, row_data in enumerate(ANSOFF['table44']):
    row = t49.rows[i + 1]
    fill_cell(row.cells[1], row_data[1], size=9, center=True, bold=True)
    fill_cell(row.cells[2], row_data[2], size=9)
    fill_cell(row.cells[3], row_data[3], size=9)

print('Аналитические таблицы заполнены')

# ===================== АБЗАЦЫ-ОТВЕТЫ =====================
A_ = C.ANSWERS
set_answer(P[54], 'Вывод:', A_['vyvod_3_1_short'])
set_answer(P[64], 'Вывод:', A_['vyvod_ge_short'])
set_answer(P[118], 'Вывод:', A_['vyvod_swot'])
set_answer(P[122], 'Выводы по матрице БКГ и сбалансированность портфеля товаров компании:',
           A_['vyvod_bcg'])
set_answer(P[224], 'ВЫВОДЫ:', A_['vyvod_cjm'])
set_answer(P[382], 'Что означают 0, 50 и 100 баллов на графике?', A_['trends_score'])
set_answer(P[413], 'Почему отчёты и карты пустые:', A_['metrika_empty'])
set_answer(P[582], 'ВЫВОД:', A_['vyvod_3_1_short'])
set_answer(P[686], 'Вывод по сегментам:', A_['vyvod_ge_short'])
for i, seg in enumerate(C.SEGMENTS):
    set_answer(P[636 + i], f'Сегмент {i+1}:', f'«{seg}»')

insert_paragraph_after(P[356], A_['mvp_choice'], size=11, align='justify')
insert_paragraph_after(P[422], A_['direct_leader'], size=11, align='justify')
insert_paragraph_after(P[437], A_['advego_vyvod'], size=11, align='justify')
insert_paragraph_after(P[439], A_['advego_lingvist'], size=11, align='justify')
insert_paragraph_after(P[441], A_['advego_iks'], size=11, align='justify')
insert_paragraph_after(P[458], A_['obshchiy_vyvod'], size=11, align='justify')
insert_paragraph_after(P[479], A_['kickoff'], size=11, align='justify')
insert_paragraph_after(P[501], A_['kp_note'], size=11, align='justify')

# второй блок критериев приёмки (Given/When/Then) для User Story № 2
h2 = insert_paragraph_after(P[349],
                            'Критерии приёмки для User Story № 2 '
                            '(«Просмотреть рейтинг надёжности поставщика»):',
                            bold=True, size=11)
tbl2_el = copy.deepcopy(T[22]._tbl)
h2._p.addnext(tbl2_el)
fill_gwt(Table(tbl2_el, doc), C.GWT[1])
h1 = insert_paragraph_after(P[345],
                            'Критерии приёмки для User Story № 1 '
                            '(«Составить график аттестации поставщиков»):',
                            bold=True, size=11)

print('Абзацы-ответы заполнены')

# ===================== ВСТАВКА ДИАГРАММ =====================
imgs = [
    (44, 'brainstorm.png', 6.4, None),
    (50, 'polygon.png', 5.4, 'Рисунок – Многоугольник конкурентоспособности'),
    (60, 'ge_mckinsey.png', 5.2, 'Рисунок – Матрица GE / McKinsey'),
    (113, 'swot.png', 6.4, 'Рисунок – SWOT-анализ компании «НОРБИТ» (доска Яндекс Концепт)'),
    (159, 'persona.png', 6.4, 'Рисунок – Портрет целевого клиента B2B (Яндекс Концепт)'),
    (209, 'cjm.png', 6.6, 'Рисунок – Карта пути клиента (Customer Journey Map)'),
    (317, 'impact_map.png', 6.6, 'Рисунок – Карта влияния Impact Mapping'),
    (355, 'usecase.png', 6.0, 'Рисунок – UML-диаграмма вариантов использования'),
    (374, 'wordstat.png', 6.2, 'Рисунок – Подбор ключевых фраз в Wordstat'),
    (381, 'trends.png', 6.2, 'Рисунок – Динамика популярности запросов в Google Trends'),
    (408, 'metrika.png', 6.2, 'Рисунок – Счётчик и сводка Яндекс.Метрики'),
    (430, 'radar.png', 6.4, 'Рисунок – Отчёты Яндекс.Радар'),
    (437, 'advego.png', 5.6, 'Рисунок – Семантический анализ текста в Advego'),
    (451, 'pagespeed.png', 5.8, 'Рисунок – Анализ сайта в Google PageSpeed Insights'),
    (507, 'kp.png', 4.7, 'Рисунок – Коммерческое предложение (страница, сделана в конструкторе)'),
    (579, 'polygon.png', 5.4, 'Рисунок – Многоугольник конкурентоспособности (раздел 3.1)'),
    (682, 'ge_mckinsey.png', 5.2, 'Рисунок – Матрица GE / McKinsey (раздел 3.2)'),
    (809, 'bcg.png', 6.3, 'Рисунок – Пузырьковая диаграмма матрицы BCG (раздел 3.3)'),
    (860, 'ansoff.png', 6.0, 'Рисунок – Матрица Ансоффа (раздел 3.4)'),
]
for pidx, img, wdt, cap in imgs:
    insert_image_after(P[pidx], A + img, width=wdt, caption=cap)
# матрица «Важность/Эффективность» в разделе 3.1 — после многоугольника
insert_image_after(P[579], A + 'importance_effectiveness.png', width=6.2,
                   caption='Рисунок – Матрица «Важность / Эффективность» (раздел 3.1)')

print('Диаграммы вставлены')

doc.save(OUT)
print('Документ сохранён:', OUT)
